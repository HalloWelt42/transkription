"""Exporte eines Transkripts: Text, Markdown, Untertitel (SRT, WebVTT), JSON und DOCX. Aufeinanderfolgende Segmente
desselben Sprechers werden in Text, Markdown und DOCX zu einem Absatz; Untertitel bleiben je Segment."""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Literal

from .modell.transkript import Segment, Transkript

Format = Literal["txt", "md", "srt", "vtt", "json", "docx"]
FORMATE: dict[str, tuple[str, str]] = {
    "txt": ("text/plain; charset=utf-8", ".txt"),
    "md": ("text/markdown; charset=utf-8", ".md"),
    "srt": ("application/x-subrip; charset=utf-8", ".srt"),
    "vtt": ("text/vtt; charset=utf-8", ".vtt"),
    "json": ("application/json", ".json"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
}


def marke(sekunden: float) -> str:
    """Zeitmarke zum Lesen: M:SS, ab einer Stunde H:MM:SS."""
    s = int(sekunden)
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60}:{s % 60:02d}"


def _untertitelzeit(sekunden: float, trenner: str) -> str:
    ms = int(round(sekunden * 1000))
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d}{trenner}{ms % 1000:03d}"


def absaetze(t: Transkript) -> Iterator[tuple[str | None, float, str]]:
    """(Sprechername, Beginn, Text) je Absatz: Segmente desselben Sprechers hintereinander werden einer."""
    sprecher, beginn, teile = None, 0.0, []
    for seg in t.segmente:
        name = t.sprecher_name(seg.sprecher)
        if teile and name != sprecher:
            yield sprecher, beginn, " ".join(teile)
            teile = []
        if not teile:
            sprecher, beginn = name, seg.start
        teile.append(seg.text)
    if teile:
        yield sprecher, beginn, " ".join(teile)


def _zeile_mit_sprecher(t: Transkript, seg: Segment) -> str:
    name = t.sprecher_name(seg.sprecher)
    return f"{name}: {seg.text}" if name else seg.text


def als_text(t: Transkript) -> str:
    mit_sprechern = bool(t.sprecher)
    return "\n\n".join(f"{n}: {text}" if mit_sprechern and n else text for n, _, text in absaetze(t)) + "\n"


def als_markdown(t: Transkript) -> str:
    zeilen = [f"# {t.name}", ""]
    marker = list(t.marker)
    for name, beginn, text in absaetze(t):
        while marker and marker[0].zeit <= beginn:
            m = marker.pop(0)
            zeilen += [f"> [{marke(m.zeit)}] {m.titel}", ""]
        zeilen += [f"**{name}** [{marke(beginn)}]:" if name else f"[{marke(beginn)}]", text, ""]
    for m in marker:
        zeilen += [f"> [{marke(m.zeit)}] {m.titel}", ""]
    return "\n".join(zeilen)


def als_srt(t: Transkript) -> str:
    bloecke = [f"{n}\n{_untertitelzeit(s.start, ',')} --> {_untertitelzeit(s.ende, ',')}\n{_zeile_mit_sprecher(t, s)}\n"
               for n, s in enumerate(t.segmente, start=1)]
    return "\n".join(bloecke)


def als_vtt(t: Transkript) -> str:
    bloecke = ["WEBVTT\n"]
    for s in t.segmente:
        name = t.sprecher_name(s.sprecher)
        text = f"<v {name}>{s.text}</v>" if name else s.text
        bloecke.append(f"{_untertitelzeit(s.start, '.')} --> {_untertitelzeit(s.ende, '.')}\n{text}\n")
    return "\n".join(bloecke)


def als_docx(t: Transkript) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    dok = Document()
    dok.add_heading(t.name, level=1)
    for name, beginn, text in absaetze(t):
        absatz = dok.add_paragraph()
        zeit = absatz.add_run(f"[{marke(beginn)}] ")
        zeit.font.color.rgb = RGBColor(0x7C, 0x8A, 0x9B)
        zeit.font.size = Pt(9)
        if name:
            absatz.add_run(f"{name}: ").bold = True
        absatz.add_run(text)
    puffer = io.BytesIO()
    dok.save(puffer)
    return puffer.getvalue()


def erzeugen(t: Transkript, format: Format) -> bytes:
    if format == "docx":
        return als_docx(t)
    if format == "json":   # ohne Stimmabdrücke: die gehören zur Erkennung, nicht zum Inhalt
        return t.model_dump_json(indent=1, exclude_none=True, exclude={"sprecher": {"__all__": {"abdruck"}}}).encode("utf-8")
    text = {"txt": als_text, "md": als_markdown, "srt": als_srt, "vtt": als_vtt}[format](t)
    return text.encode("utf-8")
